from pathlib import Path
from docx import Document
from docx.enum.section import WD_ORIENTATION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path('/home/stardust/Documents/Client/2025/INDOTIX/2026/indotix_laravel')
OUT_DIR = BASE / 'docs' / 'web-user-guide'
SHOT_DIR = OUT_DIR / 'screenshots'
OUT_FILE = OUT_DIR / 'Panduan_Penggunaan_Web_INDOTIX.docx'
LOGO = BASE / 'public' / 'logo.png'


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def set_document_style(doc: Document):
    section = doc.sections[0]
    section.orientation = WD_ORIENTATION.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)

    styles = doc.styles
    styles['Normal'].font.name = 'Calibri'
    styles['Normal'].font.size = Pt(10.5)
    styles['Title'].font.name = 'Calibri'
    styles['Title'].font.size = Pt(22)
    styles['Title'].font.bold = True
    styles['Heading 1'].font.name = 'Calibri'
    styles['Heading 1'].font.size = Pt(16)
    styles['Heading 1'].font.bold = True
    styles['Heading 2'].font.name = 'Calibri'
    styles['Heading 2'].font.size = Pt(13)
    styles['Heading 2'].font.bold = True


def add_caption(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(9)


def add_picture(doc: Document, image_path: Path, caption: str, width=Inches(9.2)):
    doc.add_picture(str(image_path), width=width)
    add_caption(doc, caption)


def add_bullets(doc: Document, items):
    for item in items:
        doc.add_paragraph(item, style='List Bullet')


def add_numbered(doc: Document, items):
    for item in items:
        doc.add_paragraph(item, style='List Number')


def add_route_table(doc: Document):
    doc.add_heading('Akses Halaman Utama', level=1)
    doc.add_paragraph(
        'Tabel berikut merangkum halaman web publik dan halaman akun pengguna yang terverifikasi dari route aplikasi.'
    )

    rows = [
        ('Beranda', '/'),
        ('Login', '/login'),
        ('Registrasi', '/register'),
        ('Event', '/events'),
        ('Hotel', '/stay'),
        ('Wisata', '/wisata'),
        ('Academy', '/academy'),
        ('Spesial Program', '/special-programs'),
        ('Retail Shop', '/retail-shop'),
        ('Keranjang Retail', '/retail-shop/cart'),
        ('Profil Pengguna', '/settings/profile'),
    ]

    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    hdr[0].text = 'Halaman'
    hdr[1].text = 'Path'
    for cell in hdr:
        set_cell_shading(cell, 'D9EDF7')
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    for label, path in rows:
        row = table.add_row().cells
        row[0].text = label
        row[1].text = path
        for cell in row:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def main():
    doc = Document()
    set_document_style(doc)

    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO), width=Inches(1.7))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.style = 'Title'
    title.add_run('Panduan Penggunaan Aplikasi Web INDOTIX')

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run('Panduan area publik dan akun pengguna web\n').bold = True
    sub.add_run('Dokumen ini disusun berdasarkan route Laravel, page component Inertia React, dan screenshot aktual aplikasi yang diverifikasi pada 24 April 2026.')

    doc.add_paragraph()
    add_bullets(doc, [
        'Dokumen ini berfokus pada penggunaan web dari sisi pengguna akhir.',
        'Fitur admin, mitra, dan proses pembayaran eksternal tidak dibahas di dokumen ini.',
        'Semua gambar di dalam dokumen diambil langsung dari tampilan aplikasi yang dijalankan saat verifikasi.'
    ])

    add_route_table(doc)

    doc.add_heading('Ringkasan Fungsi Web INDOTIX', level=1)
    doc.add_paragraph(
        'Berdasarkan navigasi publik dan route yang terverifikasi, aplikasi web INDOTIX menyediakan akses ke kategori Wisata, Event, Retail Shop, Spesial Program, Academy, dan Hotel. '
        'Pengguna juga dapat membuat akun, login, mengelola profil, serta melakukan belanja retail melalui keranjang dan checkout.'
    )

    doc.add_heading('Persiapan Sebelum Menggunakan Aplikasi', level=1)
    add_bullets(doc, [
        'Gunakan browser desktop atau laptop agar seluruh area halaman mudah terlihat.',
        'Siapkan akun pengguna jika ingin mengakses profil, keranjang, checkout, riwayat, atau halaman yang memerlukan login.',
        'Untuk checkout Retail Shop, pastikan nomor HP dan alamat utama sudah diisi pada halaman profil karena halaman checkout membaca data tersebut dari profil pengguna.'
    ])

    doc.add_heading('1. Membuka Beranda dan Navigasi Utama', level=1)
    doc.add_paragraph(
        'Beranda adalah titik masuk utama untuk mulai menjelajah produk. Dari halaman ini pengguna dapat berpindah ke kategori utama melalui menu navigasi dan melanjutkan ke area promo atau rekomendasi yang tersedia di halaman.'
    )
    add_numbered(doc, [
        'Buka alamat web INDOTIX melalui browser.',
        'Perhatikan menu kategori pada area navigasi untuk berpindah ke Wisata, Event, Retail Shop, Spesial Program, Academy, atau Hotel.',
        'Gulir ke bawah untuk melihat kartu produk atau promosi yang ditampilkan di beranda.'
    ])
    add_picture(doc, SHOT_DIR / '01-beranda.png', 'Gambar 1. Tampilan beranda web INDOTIX')

    doc.add_heading('2. Membuat Akun Pengguna', level=1)
    doc.add_paragraph(
        'Halaman registrasi menyediakan dua pilihan jenis akun, yaitu User dan Mitra. Untuk panduan penggunaan web dari sisi pembeli/pengguna akhir, langkah berikut menggunakan tipe akun User.'
    )
    add_numbered(doc, [
        'Buka halaman Daftar.',
        'Pastikan pilihan jenis akun berada pada mode User.',
        'Isi Nama Lengkap, Email Aktif, Password, dan Konfirmasi Password.',
        'Jika diperlukan, pendaftaran juga dapat dilakukan melalui tombol Daftar dengan Google.',
        'Klik tombol Buat akun untuk mengirim formulir pendaftaran.'
    ])
    doc.add_paragraph(
        'Catatan: Saat memilih tipe akun Mitra, formulir menampilkan field tambahan seperti nama usaha/event dan nomor WhatsApp. Bagian tersebut tidak dibahas lebih jauh di dokumen ini.'
    )
    add_picture(doc, SHOT_DIR / '03-register.png', 'Gambar 2. Halaman registrasi akun')

    doc.add_heading('3. Login ke Akun Pengguna', level=1)
    doc.add_paragraph(
        'Halaman login menyediakan opsi masuk menggunakan email dan password, serta tombol login Google.'
    )
    add_numbered(doc, [
        'Buka halaman Login.',
        'Masukkan email yang terdaftar pada field Email terdaftar.',
        'Masukkan password pada field Password.',
        'Jika perlu, aktifkan opsi Ingat saya.',
        'Klik Masuk sekarang untuk masuk ke akun.',
        'Jika lupa password, gunakan tautan Lupa password? yang tersedia di area form.'
    ])
    add_picture(doc, SHOT_DIR / '02-login.png', 'Gambar 3. Halaman login akun pengguna')

    doc.add_heading('4. Menjelajah Produk dari Halaman Kategori', level=1)
    doc.add_paragraph(
        'Halaman kategori dipakai untuk menemukan produk melalui pencarian, filter, pemilihan tanggal, jumlah tamu/tiket, dan pengurutan hasil. Setiap kategori memiliki karakter input yang berbeda sesuai produk yang ditawarkan.'
    )
    add_bullets(doc, [
        'Event: mencari berdasarkan nama event atau kota, memilih tanggal event, jumlah tiket, dan urutan hasil.',
        'Hotel: mencari hotel atau destinasi, memilih tanggal check-in/check-out, serta mengatur jumlah tamu dan kamar.',
        'Kategori lain seperti Wisata, Academy, Spesial Program, dan Retail Shop dapat diakses dari menu yang sama pada navigasi publik.'
    ])
    add_numbered(doc, [
        'Buka salah satu kategori dari menu utama.',
        'Gunakan kolom pencarian untuk memasukkan kata kunci yang relevan.',
        'Gunakan filter atau urutan hasil yang tersedia pada panel pencarian.',
        'Klik tombol Cari atau buka salah satu kartu produk yang muncul di daftar hasil.'
    ])
    add_picture(doc, SHOT_DIR / '04-events-list.png', 'Gambar 4. Halaman daftar Event dengan panel pencarian dan discovery')
    add_picture(doc, SHOT_DIR / '05-hotel-date-picker.png', 'Gambar 5. Halaman Hotel dengan date picker dan rekomendasi produk per tanggal')

    doc.add_heading('5. Membuka Detail Produk', level=1)
    doc.add_paragraph(
        'Setelah memilih salah satu item dari daftar, pengguna akan masuk ke halaman detail produk. Contoh yang ditampilkan di bawah berasal dari halaman detail Retail Shop karena halaman ini memperlihatkan elemen detail yang paling lengkap, seperti galeri gambar, harga, varian, kuantitas, tombol keranjang, dan tombol chat admin.'
    )
    add_numbered(doc, [
        'Dari halaman daftar, klik salah satu kartu produk.',
        'Periksa informasi utama seperti nama produk, kategori, deskripsi, harga, dan stok.',
        'Jika produk mempunyai varian, pilih salah satu varian yang tersedia.',
        'Atur jumlah pembelian pada area kuantitas.',
        'Klik Tambahkan ke Keranjang jika ingin menyimpan produk untuk dibayar kemudian.'
    ])
    add_picture(doc, SHOT_DIR / '06-retail-detail.png', 'Gambar 6. Halaman detail produk Retail Shop')

    doc.add_heading('6. Mengelola Keranjang Retail Shop', level=1)
    doc.add_paragraph(
        'Keranjang Retail Shop digunakan untuk meninjau produk yang sudah dipilih sebelum masuk ke proses checkout.'
    )
    add_numbered(doc, [
        'Pastikan sudah login ke akun pengguna.',
        'Tambahkan produk dari halaman detail ke keranjang.',
        'Buka halaman Keranjang Retail Shop.',
        'Gunakan tombol plus dan minus untuk mengubah jumlah pembelian.',
        'Gunakan tombol hapus untuk mengeluarkan item dari keranjang jika tidak jadi dibeli.',
        'Periksa subtotal dan total pada panel ringkasan.',
        'Klik Lanjutkan Pembayaran untuk masuk ke halaman checkout retail.'
    ])
    add_picture(doc, SHOT_DIR / '08-retail-cart.png', 'Gambar 7. Halaman keranjang Retail Shop')

    doc.add_heading('7. Checkout Retail Shop', level=1)
    doc.add_paragraph(
        'Halaman checkout retail dipakai untuk memeriksa data penerima dan ringkasan pesanan sebelum proses pembayaran dijalankan.'
    )
    add_numbered(doc, [
        'Buka halaman checkout dari keranjang retail.',
        'Periksa Nama Lengkap, Email, Nomor HP, dan Alamat Pengiriman.',
        'Isi catatan tambahan bila diperlukan.',
        'Pastikan ringkasan belanja sudah sesuai.',
        'Klik Lanjutkan Pembayaran untuk melanjutkan proses pembayaran.'
    ])
    doc.add_paragraph(
        'Catatan penting: pada halaman ini Nomor HP dan Alamat Pengiriman dibaca dari profil pengguna. Jika salah satu data belum ada, halaman checkout menampilkan pengingat untuk melengkapi profil terlebih dahulu.'
    )
    add_picture(doc, SHOT_DIR / '09-retail-checkout.png', 'Gambar 8. Halaman checkout Retail Shop')

    doc.add_heading('8. Mengelola Profil Pengguna', level=1)
    doc.add_paragraph(
        'Halaman profil dipakai untuk mengelola data dasar akun dan alamat pengiriman. Area ini penting terutama untuk kebutuhan checkout retail.'
    )
    add_numbered(doc, [
        'Masuk ke halaman Profil Pengguna.',
        'Perbarui data seperti Nama Lengkap, Email, dan Nomor HP pada formulir profil.',
        'Lengkapi alamat pengiriman pada bagian alamat, lalu tandai salah satu alamat sebagai alamat utama.',
        'Simpan perubahan setelah data diperbarui.'
    ])
    add_picture(doc, SHOT_DIR / '07-profile.png', 'Gambar 9. Halaman profil pengguna beserta pengelolaan alamat', width=Inches(9.0))

    doc.add_heading('Catatan Penggunaan', level=1)
    add_bullets(doc, [
        'Panduan ini hanya membahas area web yang terverifikasi secara visual dan fungsional pada saat penyusunan.',
        'Kategori utama yang tersedia di navigasi publik adalah Wisata, Event, Retail Shop, Spesial Program, Academy, dan Hotel.',
        'Dokumen ini tidak membahas panel admin, panel mitra, ataupun konfigurasi pembayaran pihak ketiga.',
        'Jika ingin mendokumentasikan alur booking per kategori secara lebih rinci, sebaiknya dibuat dokumen lanjutan per modul agar setiap proses dapat diverifikasi dengan data transaksi yang sesuai.'
    ])

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT_FILE))
    print(OUT_FILE)


if __name__ == '__main__':
    main()
